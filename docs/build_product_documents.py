from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


ROOT = Path(r"C:\RemoteDesktopControllerPlan\remote-control-mvp")
OUT = ROOT / "docs" / "product_planning"
OUT.mkdir(parents=True, exist_ok=True)

SOURCES = [
    ("S1", "MDN getDisplayMedia", "https://developer.mozilla.org/en-US/docs/Web/API/MediaDevices/getDisplayMedia", "Browser screen capture prompts the user, cannot persist permission, and requires transient activation."),
    ("S2", "Parsec technology", "https://parsec.app/technology", "Low latency comes from native control, hardware encode/decode, zero-copy GPU pipeline, frame timing, and NAT traversal."),
    ("S3", "RustDesk GitHub", "https://github.com/rustdesk/rustdesk", "Open-source remote desktop can use vendor rendezvous/relay, self-hosted relay, or custom rendezvous/relay infrastructure."),
    ("S4", "Cloudflare Realtime TURN FAQ", "https://developers.cloudflare.com/realtime/turn/faq/", "Cloudflare TURN is priced at $0.05/GB with a 1,000 GB free tier and free STUN."),
    ("S5", "Twilio Network Traversal pricing", "https://www.twilio.com/en-us/stun-turn/pricing", "Twilio STUN is free and TURN is regionally priced, starting around $0.400/GB in US/EU regions."),
    ("S6", "Tailscale free plans", "https://tailscale.com/docs/account/manage-plans/free-plans-discounts", "Tailscale Personal permits 6 free users in one tailnet and has a GitHub community option for open-source projects."),
    ("S7", "Apple App Review Guidelines", "https://developer.apple.com/app-store/review/guidelines/", "Remote desktop clients need to be generic mirrors of user-owned host devices and follow store/privacy rules."),
]

SOURCE_APPENDIX_BULLETS = [
    "For browser capture work, use MDN as the constraint source. If a proposed fix depends on the browser silently choosing the whole screen, remembering permission forever, or launching capture without user activation, the worker must label that as browser-limited and route the finished paid-product plan toward native host capture.",
    "For low-latency work, use Parsec as the performance comparison point. The project does not need to clone Parsec, but any claim about a zero-delay system should identify where the current pipeline still differs: native capture, hardware encode, hardware decode, frame pacing, UDP transport quality, jitter buffer behavior, and GPU copy count.",
    "For open-source/self-host strategy, use RustDesk as evidence that rendezvous and relay architecture can be made user-controlled. This does not mean copying RustDesk; it means the architecture should keep signaling, relay, and host trust separable so the free product can remain credible and the paid product can scale.",
    "For TURN economics, update Cloudflare and Twilio pricing before any public subscription promise. The worker should calculate relay GB per hour at multiple quality levels and should record the percentage of sessions expected to be direct P2P versus relayed.",
    "For Tailscale, keep the guidance narrow: it is a useful free workaround for technical users and private family setups, not a replacement for the product's paid remote-access onboarding. Do not make a normal user install a VPN just to understand the app.",
    "For Apple review, treat the app as a remote desktop client for user-owned host computers. Avoid language or features that make it look like hosted app streaming, a store inside a store, hidden surveillance, or unauthorized device control.",
]

SOURCE_APPENDIX_VERIFICATION_BULLETS = [
    "Do not rely on this document alone for final legal, pricing, or platform decisions.",
    "Do not use second-hand blog posts as the only authority when a primary source is available.",
    "If a source changes, update the product plan and the acceptance gate that depended on it.",
]

SOURCE_APPENDIX_ISSUE_BULLETS = [
    "Create one GitHub issue for browser-capture limitations and link it to the native host capture spike. The issue should explicitly say which behaviors are impossible or unreliable in a normal browser and which behaviors can remain in the free local fallback.",
    "Create one GitHub issue for low-latency measurement. It should require a test video, moving cursor path, input RTT log, WebRTC stats export, and a before/after comparison for each capture or rendering optimization.",
    "Create one GitHub issue for relay economics. It should include direct-vs-relay percentage, bandwidth per hour, expected free-tier usage, paid-tier fair use, provider pricing checked date, and the point where self-hosted relay becomes cheaper.",
    "Create one GitHub issue for mobile store compliance. It should collect iOS and Android privacy labels, account deletion requirements, host-device ownership language, subscription rules, support/demo mode, and review notes.",
    "Create one GitHub issue for free local/open-source positioning. It should decide which parts can be public, which assets stay private, how family-test packages are built, and how local mode stays useful even if paid services are down.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=8.4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)


def set_doc_styles(doc, title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.22)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8.7)
    normal.paragraph_format.space_after = Pt(2.2)
    normal.paragraph_format.line_spacing = 1.02
    for style_name, size, color in [
        ("Heading 1", 13.2, "1F4D78"),
        ("Heading 2", 10.2, "2E74B5"),
        ("Heading 3", 9.2, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(title).font.size = Pt(7.5)


def p(doc, text, style=None, bold_label=None):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(2.2)
    para.paragraph_format.line_spacing = 1.02
    if bold_label:
        run = para.add_run(bold_label)
        run.bold = True
        run.font.size = Pt(8.7)
        para.add_run(" " + text).font.size = Pt(8.7)
    else:
        para.add_run(text).font.size = Pt(8.7)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.left_indent = Inches(0.18)
    para.paragraph_format.first_line_indent = Inches(-0.1)
    para.paragraph_format.space_after = Pt(1.7)
    para.paragraph_format.line_spacing = 1.0
    para.add_run(text).font.size = Pt(8.45)
    return para


def add_title(doc, title, subtitle):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string("0B2545")
    p(doc, subtitle)


def add_goal_box(doc, goals):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(7.2)
    set_cell_shading(table.cell(0, 0), "E8EEF5")
    set_cell_text(table.cell(0, 0), "Completion goals for this page", bold=True, size=8.4)
    for goal in goals:
        para = table.cell(0, 0).add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(1)
        para.add_run(goal).font.size = Pt(8.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_source_table(doc):
    doc.add_heading("Research Sources Used", level=2)
    p(doc, "Use these sources as engineering constraints, not decorative citations. Future workers should re-check pricing, platform rules, and browser behavior before launch because those items can change. The sources below explain why browser capture cannot be treated as fully unattended, why native encode/decode matters for speed, why relay bandwidth must be priced carefully, and why app-store positioning must be precise.")
    for item in [
        "Before changing capture behavior, re-read the browser capture source and write down whether the change depends on a user prompt, persistent permission, or transient activation.",
        "Before promising zero-delay remote access, compare the planned media path against native low-latency products and measure capture, encode, network, decode, and draw separately.",
        "Before launching paid remote access, update TURN/STUN pricing and model bandwidth at idle, 720p30, 1080p30, 1080p60, and relay-heavy failure cases.",
        "Before submitting mobile apps, re-check Apple and Google policy language and make sure the app is positioned as user-owned-device remote control with clear privacy disclosures.",
    ]:
        bullet(doc, item)
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [0.45, 1.45, 3.0, 2.25]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    headers = ["ID", "Source", "URL", "What it proves"]
    for i, header in enumerate(headers):
        set_cell_text(table.cell(0, i), header, True, 7.5)
        set_cell_shading(table.cell(0, i), "E8EEF5")
    for sid, name, url, proof in SOURCES:
        cells = table.add_row().cells
        for i, text in enumerate([sid, name, url, proof]):
            set_cell_text(cells[i], text, size=7.1)


def add_source_appendix_detail(doc):
    doc.add_heading("Source-to-Workstream Usage Rules", level=2)
    p(doc, "This appendix page exists so the citations become actionable build rules. A future worker should not simply paste these links into a PR; they should use them to decide what is technically possible, what must be measured, and what must be disclosed to users.")
    for item in SOURCE_APPENDIX_BULLETS:
        bullet(doc, item)
    doc.add_heading("Verification Rule for Future Research", level=2)
    p(doc, "Pricing, browser behavior, app-store rules, and mobile OS capabilities can change. Before the app is shipped, submitted, priced, or marketed, the responsible worker must reopen the relevant primary source, record the date checked, and update the issue with a short statement of what changed or what remained the same.")
    for item in SOURCE_APPENDIX_VERIFICATION_BULLETS:
        bullet(doc, item)
    doc.add_heading("Issue Conversion Rules", level=2)
    p(doc, "The source appendix is complete only when it becomes work. The first GitHub migration should convert these source constraints into concrete issues so nobody has to rediscover the same limits during launch pressure.")
    for item in SOURCE_APPENDIX_ISSUE_BULLETS:
        bullet(doc, item)


def add_plan_page(doc, page):
    doc.add_heading(f"Page {page['number']}: {page['title']}", level=1)
    p(doc, page["lead"], bold_label="Purpose:")
    for section in page["sections"]:
        doc.add_heading(section["heading"], level=2)
        p(doc, section["body"])
        for item in section["bullets"]:
            bullet(doc, item)
    for section in execution_deepening_sections(page):
        doc.add_heading(section["heading"], level=2)
        p(doc, section["body"])
        for item in section["bullets"]:
            bullet(doc, item)
    add_goal_box(doc, page["goals"])
    if page["number"] < 25:
        doc.add_page_break()


BASE_DETAIL = {
    "stabilize": [
        "Treat the current Node host, phone PWA, capture page, host console, package scripts, and tests as the baseline product rather than a disposable prototype.",
        "Freeze a checkpoint before risky work, keep `dist/OpenHostLink.zip` reproducible, and prevent accidental renames like `.disabled-by-incident-response` from reaching GitHub.",
        "Every workstream must preserve real-input safety gates, session approval, trusted-device behavior, local Wi-Fi launch, and the existing high-quality WebRTC stream until a replacement proves better.",
    ],
    "performance": [
        "Measure capture, encode, network, decode, draw, and input acknowledgement separately; a single 'laggy' label is not actionable enough for product work.",
        "Favor native hardware paths for the paid high-speed tier: desktop capture APIs, GPU encode, H.264/H.265/AV1 where supported, hardware decode on Android/iOS, and frame pacing.",
        "Keep the browser/PWA as a free and universal compatibility layer, but do not pretend it can match native latency when browser capture permission, canvas drawing, and mobile Safari limits apply.",
    ],
    "security": [
        "Keep all remote control behind explicit pairing, host approval or trusted-device renewal, short-lived tokens, origin checks, command sequence validation, and a host-side kill switch.",
        "Design the paid remote tier as account-authenticated signaling plus P2P first, TURN relay fallback second, and no permanent public inbound port on a home laptop.",
        "Log security events without leaking typed text, clipboard contents, full tokens, or private screen content.",
    ],
}


PLAN_PAGES = [
    ("Product Target and Non-Negotiables", "Define the product as a premium-feeling phone remote controller where the free edition works brilliantly on local Wi-Fi and the paid edition makes remote access reliable from another network without turning the experience into a laggy slideshow.",
     [("Product shape", "The product has three client surfaces at launch: browser/PWA, Android app, and iOS app. Desktop app is deliberately postponed, but the Windows host must become clean enough that non-technical users can install it, autostart it, approve devices, and recover from errors.", ["Free local Wi-Fi must remain useful without a subscription; paid remote access must buy convenience, relay capacity, account sync, and higher reliability rather than artificially cripple the free version.", "The phone UI must stay centered on the user's original complaint: intuitive mouse control, monitor switching, keyboard, zoom, portrait and landscape usability, and visible cursor tracking.", "The publishable target is not a generic RDP clone; it is a mobile-first controller with a laptop-touchpad mental model and remote video good enough for daily lightweight control."]),
      ("Technical guardrails", "The architecture must separate control input, video transport, pairing/security, device/account management, and packaging. Each piece needs its own tests so parallel Codex instances can work without trampling each other.", BASE_DETAIL["stabilize"]),
      ("Research signal", "The current browser path is valuable but limited: MDN documents that screen capture permission cannot be persisted and requires user activation [S1]. That makes a native host path essential for unattended paid remote sessions.", ["Parsec's public technology notes point directly at the performance bar: native control, hardware encode/decode, frame timing, and smooth 60 FPS streaming [S2].", "RustDesk proves a product can be open-source/self-hostable while using rendezvous and relay infrastructure when P2P is not enough [S3].", "Cloudflare and Twilio pricing show why relay traffic must be priced carefully; video bandwidth can destroy margins if every remote session is relayed [S4][S5]."])],
     ["Everyone can explain what free local Wi-Fi includes.", "Everyone can explain what paid remote access adds.", "The team agrees that native capture/encode work is required for high-speed remote control."]),
    ("Current Build Baseline and Stabilization", "Turn the existing working local Wi-Fi controller into a stable baseline that can be moved to GitHub and shared across machines without losing the fragile fixes already discovered.",
     [("Repository truth", "The current workspace contains the Node host, public phone UI, capture page, host console, packaging scripts, docs, output checkpoints, and tests. Before GitHub migration, the repo must contain normal filenames for `package.json` and `src/server.js`, not disabled copies.", ["Run a clean file inventory: root files, `src`, `public`, `packaging`, `scripts`, `test`, `docs`, `dist`, and `output/checkpoints`.", "Move generated heavy artifacts out of the default commit path or place them under release assets; do not commit `node_modules`, smoke extraction folders, or huge acceptance output.", "Preserve `checkpoint-one-v67-current.zip`, `checkpoint-two-v68-webrtc-current.zip`, and the latest `OpenHostLink.zip` as release/reference artifacts, not active source."]),
      ("Known working behaviors", "The handoff must state the confirmed working pieces: pairing PIN, host approval, trusted sessions, real input, touchpad relative movement, keyboard, high-quality video, capture autostart, package build, and v76 touchpad/zoom fixes.", ["Document the exact commands used to run, package, smoke test, and restart the host.", "Document that monitor switching remains intentionally unresolved until the rest of the product is stabilized.", "Document that the browser version is a compatibility layer and native apps must be measured separately."]),
      ("Stabilization rules", "No new feature should land until the baseline has a Git commit, a known branch, a reproducible package, and a smoke test result. This prevents the project from becoming a pile of untraceable fixes.", BASE_DETAIL["stabilize"])],
     ["Workspace filenames are normal and runnable.", "Baseline limitations are written down without hiding them.", "The next Codex instance can run the project from the document alone."]),
    ("GitHub Migration and Repo Hygiene", "Move the project into GitHub in a way that supports multiple Codex instances, avoids accidental artifact commits, and gives the user a rollback path.",
     [("Repository creation", "Create a private GitHub repository first, then decide later whether to open-source the free local Wi-Fi edition. The initial repo should be private because the code still contains rough packaging, experimental capture paths, and host-key assumptions.", ["Initialize Git only after `.gitignore` excludes `node_modules`, `dist`, output smoke folders, log files, `.pid`, personal host keys, and local data.", "Commit source, docs, tests, package manifests, scripts, and packaging templates.", "Tag the imported baseline as `v0.1-local-wifi-baseline-v76` or similar so future experiments can be compared cleanly."]),
      ("Branch model", "Use short-lived feature branches mapped to plan chapters. Each Codex instance owns one branch and one workstream; merges happen through PRs with a checklist.", ["Example branches: `stabilize/browser-pwa`, `mobile/android-shell`, `mobile/ios-shell`, `infra/remote-relay`, `host/native-capture`, `qa/performance-harness`.", "PRs must include test commands, screenshots for UI changes, and a rollback note.", "Use GitHub Issues as the single list of product tasks, not chat fragments."]),
      ("Secrets and data", "Never commit host keys, pairing tokens, local device trust stores, logs with IP addresses, or screen-capture evidence containing personal content.", ["Add a `docs/security/dev-secrets.md` explaining where local secrets live.", "Add sample env files only with fake values.", "Add a pre-commit or CI grep for obvious secrets once the repo is stable."])],
     ["A GitHub-ready file inclusion/exclusion policy exists.", "Parallel branch names are defined.", "Release artifacts are separated from source."]),
    ("Handoff to Other Codex Instances", "Prepare the project so new Codex instances can contribute without needing the entire chat history or guessing what was tested.",
     [("Handoff packet", "Each new instance receives the handoff document, this 25-page plan, the current repo URL, the branch it owns, the exact task boundaries, and the commands it is allowed to run.", ["Include current architecture, known issues, user preferences, setup commands, testing commands, and product priorities.", "Include the rule: do not rewrite shared protocol, auth, or package scripts unless your chapter explicitly owns them.", "Include a required final report format: changed files, tests run, screenshots/artifacts, risks, and follow-up issues."]),
      ("Collision avoidance", "Split ownership by directories and interfaces. One instance can own `public/app.js` zoom UI while another owns `src/rtc-room.js`, but nobody should edit both sides of a protocol without an agreed schema issue.", ["Maintain a `docs/parallel-work-ledger.md` with active branches, owner, files, status, and merge order.", "Use interface docs for WebSocket messages, RTC signaling, settings schema, and package launch contracts.", "Keep a merge captain role that reviews PRs and resolves conflicts in a controlled order."]),
      ("Continuity prompt", "The handoff prompt must be operational, not motivational. It should say what exists, what to protect, what to ignore, how to run it, and how to report back.", ["Avoid asking new instances to 'make it better' broadly.", "Give each instance one acceptance gate and one directory boundary.", "Require them to update docs when they change behavior."])],
     ["A new Codex instance can start without this chat.", "File ownership is explicit.", "Merge/report rules are clear enough to prevent collateral damage."]),
    ("Architecture Split: Free Local vs Paid Remote", "Define the product architecture around two modes that share UX and host code but differ in network path, account services, and support expectations.",
     [("Free local edition", "The free edition should run host and phone on the same LAN, plus optional personal Tailscale for technical users. It can be open-source and GitHub-downloadable.", ["Use LAN discovery/QR/manual URL and keep Tailscale as a documented free different-network workaround for non-commercial/power users [S6].", "Do not require account login for local mode; pairing PIN and host approval are enough.", "Keep browser/PWA support because it lowers friction and supports quick testing."]),
      ("Paid remote edition", "The paid edition should provide account login, device registry, signaling, NAT traversal, relay fallback, better onboarding, persistent device pairing, and support.", ["Prefer P2P direct connections first; use TURN relay only when NAT/firewalls block direct traffic.", "Meter or cap relay bandwidth by subscription tier because 1080p60 video can burn GB quickly.", "Use managed TURN initially for launch speed, then evaluate self-hosted relay when traffic volume justifies it."]),
      ("Shared core", "Control protocol, input safety, monitor metadata, device trust, UI language, and support diagnostics should be shared across free and paid editions.", ["Avoid forking the UI into unrelated products.", "Use capability flags from the host/account service to show remote-only features.", "Keep the local edition as a credibility engine and debugging fallback."])],
     ["Free and paid boundaries are explicit.", "No artificial free-version breakage is required.", "Remote costs are designed into the subscription model."]),
    ("Host Capture and Video Pipeline", "Replace the fragile browser capture dependency with a host capture/encoder pipeline capable of smooth remote video while keeping the browser capture path as fallback.",
     [("Current limitation", "The existing high-quality stream depends on a Chrome capture window and WebRTC. It can work well, but browser screen sharing cannot be fully unattended in the general case because browser permission is user-mediated [S1].", ["Keep the current capture page for local free mode and fallback.", "Document that auto-select flags are browser-specific and may break across machines.", "Do not build a business promise on browser capture alone."]),
      ("Native capture direction", "The host should evolve toward a native capture module using Windows Graphics Capture or Desktop Duplication, then hardware encode through Media Foundation/NVENC/AMF/QuickSync where available.", ["Start with Windows because that is the current host target.", "Expose a clean stream interface so later macOS/Linux host support can be added without rewriting phone clients.", "Track capture FPS, encode latency, dropped frames, resolution, cursor composition, and monitor ID."]),
      ("Performance bar", "Parsec's public notes set the engineering target: native low-level control, zero-copy GPU path, hardware decode, and frame-timing optimization [S2]. The paid tier must move in that direction.", BASE_DETAIL["performance"])],
     ["Browser capture remains fallback, not the final paid path.", "Native capture module is specified.", "Performance metrics are defined before implementation."]),
    ("Remote Networking, NAT Traversal, and Relay", "Design remote access so most sessions are direct and fast, with relay as a paid fallback rather than the default path.",
     [("Connection setup", "Use an account-backed signaling service: host registers online status, phone requests a session, host approves/trusts device, then both sides exchange WebRTC offers, ICE candidates, and capability metadata.", ["Use STUN for direct UDP discovery.", "Use TURN only after direct candidates fail or are predicted to fail.", "Keep a WebSocket control fallback only for low-bandwidth command signaling, not video."]),
      ("Relay choices", "Cloudflare Realtime TURN is currently attractive for early launch because it has global infrastructure, $0.05/GB pricing, free STUN, and a 1,000 GB free tier [S4]. Twilio is mature but much more expensive for TURN egress in many regions [S5].", ["Prototype with Cloudflare TURN and cost dashboards.", "Add provider abstraction so TURN vendor can be swapped.", "If usage grows, evaluate self-hosted relay nodes like RustDesk-style rendezvous/relay [S3]."]),
      ("Remote safety", "Remote access increases abuse risk, so the host must show active sessions, allow one-click stop, require explicit trust for unattended use, and log remote connections clearly.", BASE_DETAIL["security"])],
     ["Remote path is P2P-first.", "TURN cost is measured and priced.", "The relay provider can be changed later."]),
    ("Browser/PWA Version Repair", "Fix the browser version as a stable free controller surface on mobile and desktop, while being honest about browser capture limitations.",
     [("Mobile browser", "The PWA must be clean on iPhone Safari, Android Chrome, portrait, landscape, notch/safe-area, and home-screen installed mode. Its controls should not rely on hover, desktop-only keyboard shortcuts, or exact viewport height.", ["Retest keyboard toggle, touchpad relative movement, tap/double-tap/hold, two-finger scroll, zoom lens, display sheet, settings sheet, and reconnect.", "Keep touchpad visual state separate from real cursor input forever; the v76 fix must become a regression test.", "Add a mobile browser compatibility matrix to docs."]),
      ("Desktop browser", "Desktop browser support should exist for testing and emergency control, but desktop users are not the primary UX. It needs a sensible layout, keyboard shortcuts, pointer lock where allowed, and no mobile-only assumptions.", ["Add responsive desktop styles.", "Use the same pairing/session model.", "Avoid exposing host secrets in desktop test pages."]),
      ("Capture limits", "Browser capture can prompt, require activation, and refuse persistent permission [S1]. The browser edition can display and control a stream, but final unattended host capture should move native.", ["Keep capture status visible.", "Explain when a user must touch the laptop.", "Use native path for paid remote promise."])],
     ["Mobile and desktop browser test matrix exists.", "Known browser capture limits are documented.", "v76 touchpad/zoom regressions stay covered."]),
    ("Android App Plan", "Build an Android app that starts as a polished WebView/native shell around the existing controller, then progressively replaces performance-sensitive pieces with native code.",
     [("Phase 1 shell", "Use Kotlin/Android or React Native/Flutter only after choosing the long-term cross-platform strategy. The quickest path is likely a native Android shell that embeds the PWA, manages device identity, push/reconnect, safe-area, haptics, and secure storage.", ["Store tokens in Android Keystore-backed storage.", "Use native haptics and keyboard integration.", "Expose share/open-link handling for local LAN URLs and paid account login."]),
      ("Phase 2 media", "Move video receive/decode into native WebRTC where needed. Browser WebView may not provide the same hardware decode and timing control as native WebRTC.", ["Use native WebRTC stats for jitter, decode, dropped frames, and RTT.", "Keep the touchpad UI native or hybrid only if gesture fidelity is perfect.", "Build an internal latency overlay."]),
      ("Play Store readiness", "Remote-control apps must present a clear user-owned host model, privacy disclosures, account deletion path, and no deceptive accessibility usage.", ["Do not request Android Accessibility permissions for controlling the phone; this app controls the remote computer.", "Declare network/device data collection accurately.", "Include a demo/practice mode for review if the host app is required."])],
     ["Android MVP scope is split into shell and native media phases.", "Secure storage and review needs are included.", "Gesture fidelity remains the product differentiator."]),
    ("iOS App Plan", "Build an iPhone app that is App Store-compliant, mobile-first, and honest about remote desktop scope.",
     [("Apple constraints", "Apple's guideline 4.2.7 allows remote desktop clients under specific conditions when mirroring user-owned host devices and not acting as a store-like interface [S7]. The iOS app must be a generic remote desktop/control client, not a way to run or buy hosted software.", ["All account management and subscriptions need to follow App Store rules when sold in-app.", "Do not mimic iOS or App Store UI inside the remote stream.", "Include privacy disclosures for device identifiers, diagnostics, and relay usage."]),
      ("MVP implementation", "Start with SwiftUI plus WKWebView only if the PWA UX is already stable, then evaluate native WebRTC and native gestures for performance.", ["Use Keychain for tokens.", "Use native haptics, safe-area handling, keyboard accessory controls, and orientation-specific layouts.", "Provide a demo mode or guided screenshots for App Review because the app depends on a host."]),
      ("Native media path", "For paid remote speed, plan native WebRTC receive/decode and a custom gesture layer. The browser path remains a bootstrap, not the final App Store experience.", BASE_DETAIL["performance"])],
     ["iOS review constraints are reflected in product design.", "MVP and native media phases are separated.", "Store/privacy requirements are not left until launch week."]),
    ("Input, Gesture, and UX Finalization", "Make the controls feel intentional rather than experimental, across all client surfaces.",
     [("Touchpad model", "The primary mobile control model should be touchpad-relative movement, single tap left click, double tap right click, hold/drag gestures, two-finger scrolling, edge-hold movement, and optional direct screen touch.", ["The visual touchpad dot must never represent the real cursor; it is a local joystick/hint only.", "Add sensitivity presets for fine, normal, fast, and presentation modes.", "Add tutorial overlays only during onboarding, not as permanent clutter."]),
      ("Zoom model", "Zoom should have two separate concepts: view zoom that crops the whole display, and cursor lens that follows the mouse. Settings must control lens size, lens zoom, hold-to-show, and auto-follow.", ["Lens cursor is a small gold dot centered in the lens.", "When auto-follow is enabled, pan toward cursor smoothly and clamp only at real screen edges.", "When zoom is disabled, remove lens and reset view cleanly."]),
      ("Keyboard and shortcuts", "The keyboard button should summon the native keyboard on mobile, while shortcut buttons handle Windows, Ctrl hold, Codex, Claude, and user-defined app launches.", ["Shortcuts belong in a configurable rail or sheet.", "Never make the rail so wide it steals landscape control space.", "All shortcuts must fail visibly if the host app cannot launch them."])],
     ["Touchpad, zoom, and keyboard are defined as separate systems.", "The visual dot cannot move the real cursor.", "Shortcut UX is configurable instead of hard-coded forever."]),
    ("Monitor Switching and Multi-Display Accuracy", "Fix monitor switching with explicit identity, source lock, and calibration instead of guessing from fragile browser source order.",
     [("Current issue", "The known unresolved problem is that switching to another monitor can briefly show the correct screen, then black out or return to screen one. This is likely source-selection/capture identity drift, especially across machines.", ["Do not build more features on top of broken monitor identity.", "Record monitor bounds, logical bounds, scale factor, display name, adapter ID where available, and capture source ID.", "Keep manual source lock as an emergency fallback but make auto-detect better."]),
      ("Native path", "Native capture should remove most browser source ambiguity by selecting a monitor through OS APIs rather than Chrome's picker labels.", ["For Windows, enumerate displays and capture item IDs.", "Map input coordinates using physical and logical bounds.", "Persist per-machine monitor calibration keyed by stable display signature."]),
      ("Testing", "Build a fake multi-monitor harness plus real physical acceptance checklist.", ["Test side-by-side, stacked, mixed DPI, portrait, identical resolution, laptop + external, and monitor disconnect/reconnect.", "Verify video source, cursor overlay, direct touch, touchpad, zoom lens, and screenshot fallback all agree.", "Add a manual 'this screen is wrong' report button that exports source diagnostics."])],
     ["Monitor switching is treated as a system, not a button bug.", "Native capture is the preferred fix path.", "Mixed-DPI/multi-monitor tests are specified."]),
    ("Security, Trust, and Abuse Prevention", "Build security as a product feature because remote desktop tools are high-risk by nature.",
     [("Threat model", "Assets include screen contents, keyboard/mouse control, account identity, device trust, relay credentials, host logs, and billing state. Attackers include wrong-phone local users, stolen tokens, malicious relay clients, and social-engineering misuse.", ["Create a formal threat model before paid remote beta.", "Separate pairing PIN, session token, trusted device credential, and account identity.", "Rotate or revoke every credential type."]),
      ("Host control", "The host must always have visible state and control: active sessions, last device, remote IP/network path, streaming status, real input toggle, release all buttons, revoke trust, and stop server.", BASE_DETAIL["security"]),
      ("Compliance posture", "Privacy policy, data inventory, data deletion, store disclosures, and abuse-report pathway must exist before app review or paid launch.", ["Never log raw typed text or screen frames by default.", "Make diagnostics opt-in for support exports.", "Keep subscription billing separate from host control permission."])],
     ["Threat model scope is explicit.", "Host-side stop/revoke controls are mandatory.", "Privacy/logging rules are not optional."]),
    ("Account, Billing, and Subscription Design", "Design the paid tier around real cost drivers and user value instead of vague subscriptions.",
     [("Subscription value", "Paid remote access should include remote signaling, TURN relay fallback, device registry, account sync, trusted device recovery, priority capture path, and support diagnostics.", ["Free local Wi-Fi stays free.", "Tailscale remains documented for technical users where licensing permits.", "Paid removes setup friction and improves reliability, not basic local control."]),
      ("Cost model", "Relay cost is the main variable. Cloudflare's $0.05/GB and 1,000 GB free tier are attractive for early testing [S4], while Twilio's regional TURN prices starting around $0.400/GB show how expensive the wrong provider can be [S5].", ["Estimate 720p30, 1080p30, 1080p60, and idle bandwidth.", "Set fair-use limits or quality caps by plan.", "Prefer direct P2P and reduce relay usage through NAT traversal quality."]),
      ("Billing implementation", "Use Stripe for web billing initially if store rules allow; for iOS digital subscription features, evaluate Apple IAP obligations carefully.", ["Keep entitlement checks server-side.", "Let users see relay usage and device count.", "Do not block emergency local access because billing failed."])],
     ["Paid value maps to costs.", "Relay usage is priced/capped.", "Billing rules are separated for web, Android, and iOS."]),
    ("Infrastructure and Backend Services", "Create the minimum backend that enables paid remote access without overbuilding enterprise infrastructure.",
     [("Services", "Initial backend services: account auth, device registry, signaling WebSocket, ephemeral TURN credential minting, subscription entitlement, telemetry ingestion, and support log upload.", ["Use a managed database and auth provider only if it reduces launch risk.", "Keep signaling stateless where possible.", "Store only metadata needed for operation and support."]),
      ("Deployment", "Start with a simple cloud deployment that supports WebSockets globally enough for beta; then add regional routing if latency data proves it is needed.", ["Separate production, staging, and local dev.", "Use IaC once the shape stabilizes.", "Add uptime checks for signaling and TURN credential endpoints."]),
      ("Observability", "Track connection setup time, direct vs relay percentage, TURN GB, input RTT, video FPS, dropped frames, session disconnect reasons, and capture errors.", ["Do not capture screen contents.", "Make support bundles user-initiated.", "Build dashboards before charging broadly."])],
     ["Backend services are scoped to paid remote access.", "Telemetry is useful but privacy-preserving.", "Deployment complexity grows only when metrics demand it."]),
    ("Performance Measurement and Optimization", "Make speed measurable so optimization work is not guesswork.",
     [("Latency budget", "Break glass-to-glass delay into capture, encode, network, jitter buffer, decode, draw, and input feedback. Set target budgets for local Wi-Fi and paid remote.", ["Local target: feel instant for mouse and text, high-quality stream at smooth FPS.", "Remote target: P2P as close to local as network allows, relay acceptable but transparent.", "Record device class because old phones decode differently."]),
      ("Optimization path", "Prioritize the largest delays first: browser capture permission and source drift, canvas decode/draw hitches, WebRTC frame pacing, mouse jitter, and relay path selection.", BASE_DETAIL["performance"]),
      ("Test harness", "Build synthetic tests and physical tests: moving clock video, scrolling page, tab dragging, cursor path smoothness, zoom lens follow, monitor switch, and low-bandwidth relay.", ["Use automated stats plus screen recordings.", "Make performance reports comparable across commits.", "Block release if FPS or input RTT regresses beyond threshold."])],
     ["Latency budget exists.", "Optimization targets are measurable.", "Release gates include performance regression checks."]),
    ("Packaging, Install, and Autostart", "Make installation boring for non-technical users while preserving a developer-friendly source build.",
     [("Free GitHub download", "Provide a GitHub release zip or installer for Windows local Wi-Fi mode. The zip should contain one obvious launcher, instructions, Node runtime strategy, and no confusing stale scripts.", ["Decide whether to bundle Node or require installation; for non-technical users, bundled runtime or single installer is better.", "Generate a signed installer later, but keep zip for beta.", "Write `START-HERE.txt` in plain language."]),
      ("Host autostart", "Remote value depends on the host being available when the user is away. Add opt-in startup registration, tray status, health check, and restart-on-crash strategy.", ["Do not silently run remote control without visible tray/state.", "Let users disable autostart easily.", "Warn if firewall or capture permission blocks remote use."]),
      ("Update path", "Add versioning, changelog, and migration rules. Do not strand trusted devices or settings when upgrading.", ["Semver app and protocol.", "Backward-compatible protocol where possible.", "One rollback package per release."])],
     ["Packaging is beginner-friendly.", "Autostart is opt-in and visible.", "Update/rollback rules exist."]),
    ("Testing, CI, and Release Gates", "Create evidence that the product works before public release and before charging users.",
     [("Automated tests", "Keep existing Node tests and add CI for protocol, input adapter, session store, RTC room, settings, browser UI, package smoke, and security lint.", ["CI must run on PRs.", "Tests that require real screen capture can be nightly/manual.", "Package smoke should use a clean extraction directory."]),
      ("Manual acceptance", "Some features need physical devices: iPhone Safari, Android Chrome, Android app, iOS app, multi-monitor Windows, cellular remote path, and weak Wi-Fi.", ["Create checklists with screenshots and expected timings.", "Record device, OS, browser/app version, network type, and result.", "Do not mark a release ready without real phone tests."]),
      ("Release decision", "A release candidate needs green tests, no critical security issues, performance within target, install instructions verified by a fresh machine, and known issues listed.", ["Publish release notes.", "Attach zip/installer hashes.", "Tag the Git commit."])],
     ["CI scope is defined.", "Physical-device gates are defined.", "Release criteria are explicit."]),
    ("UI Polish and Product Feel", "Make the interface feel premium, minimal, and reliable rather than like a diagnostics page.",
     [("Phone UI", "Keep the black/gold premium theme if it remains readable. Controls should be large enough, evenly sized, and limited to what the user needs during actual control.", ["Primary rail: Keys, Zoom, Display, Settings, Ctrl/Win/shortcuts as configured.", "Hide raw counters and debug status from normal mode.", "Landscape mode should prioritize screen size and usable controls."]),
      ("Host UI", "The host console should be calm: connection URL, QR, PIN, sessions, capture status, input safety, settings, logs/export. It should not look different between dev and package unless intentionally branded.", ["One launch path.", "Clear capture status.", "Clear remote/local network path."]),
      ("Onboarding", "Onboarding must explain: install host, open phone link or app, pair PIN, approve, enable real input, keep signed in, remote access setup.", ["Use short steps and visual QR.", "Add troubleshooting only when needed.", "Include practice/demo mode for store review and nervous users."])],
     ["Normal UI hides debug clutter.", "Host/package UI are consistent.", "Onboarding maps to real setup steps."]),
    ("Documentation and Support System", "Turn the project's scattered knowledge into durable docs that users and Codex instances can rely on.",
     [("Developer docs", "Create docs for architecture, protocol, capture lifecycle, input mapping, monitor switching, packaging, tests, and deployment.", ["Every shared interface gets a short contract file.", "Every tricky bug gets a regression note.", "Docs should name owners and acceptance gates."]),
      ("User docs", "Create quick-start, troubleshooting, remote access setup, security explanation, and uninstall guides.", ["Write for non-technical family users first.", "Keep advanced Tailscale/local firewall steps separate.", "Add screenshots after UI stabilizes."]),
      ("Support docs", "Support needs diagnostic export, version info, network path, relay/direct status, capture stats, and safe redaction.", ["No raw screen screenshots unless user explicitly attaches them.", "No raw typed text.", "Make logs understandable enough for issue reports."])],
     ["Developer and user docs are separate.", "Support export is privacy-aware.", "Bug knowledge becomes docs/tests."]),
    ("Store, Legal, and Policy Readiness", "Prepare for app review and payment rules before the apps are submitted.",
     [("iOS", "Apple's remote desktop rule requires careful positioning as a generic remote desktop client for user-owned host devices [S7]. Avoid store-like browsing inside the remote stream and prepare a demo account/mode.", ["Privacy labels must cover diagnostics, identifiers, account info, and crash logs.", "Subscription/IAP policy must be reviewed before selling remote access inside iOS.", "App review notes must explain host setup clearly."]),
      ("Android", "Prepare Play Store data safety, privacy policy, account deletion, and no unnecessary sensitive permissions.", ["Do not request microphone unless laptop-output audio architecture actually needs it and user chooses it.", "Do not misuse accessibility services.", "Use foreground service notifications if background connectivity is needed."]),
      ("Legal", "Remote desktop can be abused, so terms must forbid unauthorized control, require user-owned devices, and explain logs, relay, and security.", ["Add privacy policy.", "Add terms of service.", "Add responsible-use language."])],
     ["Store policy risks are known early.", "Privacy/legal docs are on the roadmap.", "App review demo strategy is included."]),
    ("Parallel Workstream Map", "Divide work so multiple Codex instances can move quickly without clobbering one another.",
     [("Chapter ownership", "Assign separate chapters to separate branches. The merge captain updates shared docs and resolves protocol changes.", ["Workstream A: repo hygiene, CI, GitHub migration.", "Workstream B: browser/PWA polish and regression tests.", "Workstream C: Android shell prototype.", "Workstream D: iOS shell prototype.", "Workstream E: native host capture research/spike.", "Workstream F: remote signaling/TURN backend prototype."]),
      ("Interface locks", "Shared interfaces need mini RFCs before change: WebSocket messages, RTC signaling, auth/session schema, settings storage, capture metadata, billing entitlement.", ["A worker may add fields compatibly.", "A worker may not rename/remove shared fields without RFC approval.", "Each PR updates tests and docs."]),
      ("Merge order", "Merge stabilization first, then browser fixes, then backend prototypes, then mobile shells, then native capture experiments.", ["Keep risky native capture on a branch until it beats current stream.", "Keep paid backend behind feature flags.", "Never block local free mode on paid services."])],
     ["Workstreams have owners and branch names.", "Shared interfaces are protected.", "Merge order avoids collateral damage."]),
    ("90-Day Execution Roadmap", "Sequence the work from handoff to beta without pretending everything can be done at once.",
     [("Days 1-14", "Stabilize repo, GitHub, CI, handoff docs, package, browser regression tests, and issue tracker. No paid backend work should distract from making the baseline reproducible.", ["Create private repo and baseline tag.", "Fix package naming and one-launcher confusion.", "Write architecture and protocol docs.", "Run full test suite and package smoke on a clean PC."]),
      ("Days 15-45", "Build browser/PWA polish, Android shell, iOS shell, remote signaling prototype, and native capture spike in parallel branches.", ["Measure current latency.", "Prototype Cloudflare TURN credentials.", "Prototype app shells with real pairing.", "Prototype Windows native capture feasibility."]),
      ("Days 46-90", "Choose the high-speed path, harden account/subscription backend, run closed beta, polish onboarding, and prepare store submissions.", ["Beta with 5-20 trusted testers.", "Collect relay cost data.", "Decide launch pricing.", "Prepare App Store/Play Store review assets."])],
     ["Roadmap is phased.", "Baseline stabilization comes first.", "Beta and monetization are tied to measured quality."]),
    ("Later Desktop App and Platform Expansion", "Keep the desktop client out of the immediate launch scope while still designing today's protocol, account model, and media path so a future desktop viewer is not boxed in.",
     [("Why it waits", "The user explicitly wants desktop app work later. That is the right call because the current bottleneck is host reliability, mobile control quality, remote networking, and native phone apps.", ["Do not split attention into Electron/Tauri desktop clients before mobile and paid remote access are proven.", "Keep the browser desktop view usable for testing, but do not polish it at the expense of phone UX.", "Record desktop-client assumptions now so protocol choices do not block it later."]),
      ("Protocol readiness", "A future desktop viewer should reuse the same signaling, device trust, entitlement, capture metadata, and input command model.", ["Avoid mobile-only names in shared protocol fields.", "Keep keyboard/mouse commands expressive enough for desktop clients.", "Keep stream negotiation capability-based instead of hard-coding phone constraints."]),
      ("Expansion order", "After paid remote beta proves speed and economics, evaluate desktop clients, macOS host support, Linux host support, browser extensions, and enterprise management.", ["Desktop viewer can become a paid productivity feature later.", "macOS/Linux host support require separate capture/input/security research.", "Enterprise features should wait until consumer reliability is solid."])],
     ["Desktop app is intentionally deferred.", "Current architecture remains future-compatible.", "Expansion does not distract from mobile-first launch."]),
    ("Acceptance Gates and Goal Ledger", "Summarize every goal without skipping steps so the plan can become GitHub issues and Codex chapters.",
     [("Gate 1: handoff", "The project is not ready for parallel work until the handoff document exists, source files are restored, GitHub repo is created, baseline branch/tag exists, and a new Codex instance can run tests from the docs alone.", ["Handoff doc delivered.", "Transfer plan to user's PC documented.", "Codex prompt pack included.", "Git ignore and artifact policy defined."]),
      ("Gate 2: product foundation", "The free local product is shippable when the browser/PWA, host console, packaging, tests, and install docs work on a clean Windows machine and at least one iPhone and one Android phone.", ["No debug clutter in normal UI.", "Keyboard/touchpad/zoom/display/settings pass manual checks.", "Package zip or installer is clean.", "Known monitor-switching limitation is either fixed or clearly marked before public release."]),
      ("Gate 3: paid remote", "The paid version is viable only after P2P remote sessions work, TURN fallback is measured, relay cost fits pricing, native/mobile performance beats the browser path, security review passes, and closed beta users say it feels smooth enough.", ["Cloudflare/Twilio/self-host relay decision documented.", "Subscription entitlement works.", "Store/legal requirements are satisfied.", "Performance dashboard proves remote speed."]),
      ("Sources", "Evidence used in this plan is carried into the handoff manual's research-source appendix. Future instances should use that appendix to verify browser capture limits, low-latency architecture assumptions, TURN pricing, Tailscale positioning, and app-store policy before making launch promises.", ["Do not repeat vague research; reopen the primary source when the decision is pricing, platform policy, browser behavior, or paid remote-access architecture.", "Convert each source constraint into a GitHub issue so research becomes work instead of decoration."])],
     ["All chapters have acceptance gates.", "No major product area is missing.", "Sources are attached for future verification."]),
]


def execution_deepening_sections(page):
    title = page["title"]
    return [
        {
            "heading": "Full build checklist",
            "body": f"For {title}, the worker must treat the page as a build manual. The checklist below is deliberately concrete so it can be copied into GitHub issues without losing the operational steps.",
            "bullets": [
                "Start by writing the current behavior in one paragraph using real evidence from the workspace, package, browser, phone, server log, or test output.",
                "List every file that will be read before editing and every file expected to change. If the list grows, stop and update the issue scope before continuing.",
                "Identify the user-visible behavior, the host-side behavior, the phone-side behavior, the security behavior, and the failure behavior separately.",
                "Define the fallback path before implementing the new path. The local Wi-Fi controller must remain usable even if a paid, native, mobile, or relay experiment fails.",
                "Add or update tests that lock the original bug or capability. Where automation cannot cover it, write the physical-device test with exact device, browser/app, network, display layout, and expected result.",
                "Update the documentation in the same branch, including user-facing instructions if behavior changed and developer notes if a contract, protocol, or launch command changed.",
                "Package or run from a clean checkout before calling the issue finished. A feature that only works in the warm development session is not done.",
            ],
        },
        {
            "heading": "Detailed implementation steps",
            "body": f"The implementation for {title} should move through these steps in order. Skipping the early evidence and contract work is how the prototype kept regressing.",
            "bullets": [
                "Step 1: capture baseline evidence. Record current version, command used, URL used, package name, branch, device, and the exact broken or missing behavior.",
                "Step 2: write the interface contract. For UI work, this means visible states and gestures. For host work, this means endpoints, messages, settings, or process behavior. For backend work, this means auth, session, entitlement, and network states.",
                "Step 3: implement the smallest vertical slice that proves the contract. Avoid sweeping rewrites unless the issue is explicitly a rewrite chapter.",
                "Step 4: add regression protection. Prefer unit tests for pure logic, browser tests for UI state, package smoke for launch behavior, and physical checklists for phone/capture behavior.",
                "Step 5: run the exact acceptance gate. If the gate cannot run locally, document the blocker and attach the next-best evidence instead of pretending it passed.",
                "Step 6: write the handoff note. Include what changed, what did not change, what remains risky, and how to roll back without losing unrelated work.",
            ],
        },
        {
            "heading": "Rejection criteria",
            "body": f"A reviewer should reject work on {title} when it only looks finished. These rejection cases are designed to prevent a repeat of broken screen capture, jittery controls, stale packages, and machine-specific behavior.",
            "bullets": [
                "Reject if the worker cannot name the test command, package command, or physical-device check that proves the change.",
                "Reject if a UI change hides a bug by removing a control, hiding debug state, or reducing functionality without documenting the tradeoff.",
                "Reject if a speed improvement is claimed without FPS, latency, dropped-frame, RTT, relay/direct, or device evidence.",
                "Reject if the change works on the development laptop but the branch has no plan for fresh Windows machines, different Chrome installs, different monitor layouts, or phone browsers.",
                "Reject if a remote-access or trusted-device change weakens pairing, host approval, token expiry, revoke, release-buttons, or stop-control behavior.",
                "Reject if source, package, docs, and release artifact disagree about the launcher name, host URL, QR code behavior, or setup steps.",
            ],
        },
        {
            "heading": "Execution sequence",
            "body": f"This page becomes work only when it is converted into issues with owners, files, tests, and acceptance evidence. For {title}, the sequence below is the default order unless a later spike proves a dependency is wrong.",
            "bullets": [
                f"Open or update a GitHub issue named after '{title}', link this page, and list the exact files or services that are in scope.",
                "Create a short branch from the current baseline; do not mix unrelated cleanup, UI polish, backend experiments, and package changes in one branch.",
                "Write or update the contract first: expected behavior, data shape, command, UI state, or user-facing promise.",
                "Implement the smallest complete slice that can be tested end to end, then add regression coverage before expanding the slice.",
                "Finish with proof: command output, package smoke, browser/mobile screenshot, performance numbers, or physical-device acceptance notes as appropriate.",
            ],
        },
        {
            "heading": "Parallel ownership and collision rules",
            "body": f"Parallel Codex instances may work near {title}, but only if each owns a different boundary and updates the work ledger before editing.",
            "bullets": [
                "UI-only work owns `public/*` plus matching tests; protocol or server behavior changes require a separate interface note.",
                "Host/backend work owns `src/*` or cloud service code and must preserve local free mode unless the issue explicitly changes that contract.",
                "Mobile app work must not silently fork control semantics; any gesture or command difference needs a compatibility note.",
                "Packaging/docs work can update instructions and release artifacts, but must not hide known limitations or mark unverified behavior as done.",
            ],
        },
        {
            "heading": "Failure modes to check before closing",
            "body": f"The common failure pattern for this project is a feature looking good on one machine while breaking another tester's phone or monitor setup. Closing {title} requires checking the failure modes, not just the happy path.",
            "bullets": [
                "Fresh machine or clean profile behaves differently from the development laptop.",
                "Phone Safari, Android Chrome, or app WebView handles touch, keyboard, screen size, or media playback differently.",
                "Multi-monitor, mixed-DPI, external display, or reconnect path changes the coordinate/capture assumptions.",
                "A paid-remote change accidentally breaks the free local Wi-Fi path or requires an account where local mode should not.",
                "A performance fix improves FPS while making input jitter, zoom follow, audio, or battery usage worse.",
            ],
        },
        {
            "heading": "Issue breakdown for a worker",
            "body": f"Convert {title} into a set of small issues instead of one huge vague assignment. Each issue should have a single mergeable result and a clear rollback point.",
            "bullets": [
                "Issue 1: write the exact current-state evidence before changing anything, including file names, commands, screenshots, or API output.",
                "Issue 2: define the target behavior in one short contract so future code, UI, docs, and tests agree.",
                "Issue 3: implement the first narrow path and keep the old path available behind a flag or fallback until the new path is proven.",
                "Issue 4: add regression coverage for the exact bug or capability, including the machine/device condition that originally exposed it.",
                "Issue 5: update the user-facing documentation and the developer handoff notes so the next instance does not rediscover the same detail.",
                "Issue 6: package or deploy the change only after the branch can be tested from a clean checkout or clean extraction.",
            ],
        },
        {
            "heading": "Evidence artifacts to attach",
            "body": f"The work for {title} is not complete just because the code looks plausible. Attach proof that a reviewer can inspect without trusting the worker's memory.",
            "bullets": [
                "Command transcript or test output showing the relevant automated tests passed.",
                "Screenshots or rendered page previews for any user-interface, onboarding, store, or document change.",
                "Performance numbers for any capture, networking, video, input, relay, or mobile-app claim.",
                "A short risk note naming what was not tested, especially physical phones, different networks, and multi-monitor setups.",
                "A rollback note naming the branch, tag, zip, feature flag, or previous implementation to return to if testers report a regression.",
            ],
        },
        {
            "heading": "Handoff notes for the next instance",
            "body": f"When another Codex instance picks up {title}, it should start from these notes rather than from the chat history.",
            "bullets": [
                "Read this page, the handoff manual, and the files named by the issue before editing.",
                "State the file ownership boundary in the first progress update.",
                "Do not mark the issue complete unless the acceptance evidence is stronger than a visual guess.",
                "If a blocker appears, leave a precise reproduction path and continue with adjacent documentation or tests instead of stopping the whole product plan.",
            ],
        },
    ]


def expand_plan_pages():
    pages = []
    for idx, (title, lead, sections, goals) in enumerate(PLAN_PAGES, start=1):
        normalized_sections = []
        for heading, body, bullets in sections:
            normalized_sections.append({"heading": heading, "body": body, "bullets": bullets})
        pages.append({"number": idx, "title": title, "lead": lead, "sections": normalized_sections, "goals": goals})
    return pages


def create_plan_doc():
    doc = Document()
    set_doc_styles(doc, "Remote Controller Product Execution Plan")
    add_title(doc, "Remote Controller Product Execution Plan", "A 25-page, chapterized build plan for turning the current local Wi-Fi controller into a GitHub-ready free product plus paid high-speed remote access.")
    for page in expand_plan_pages():
        add_plan_page(doc, page)
    path = OUT / "Remote_Controller_Product_Execution_Plan_25_Pages.docx"
    doc.save(path)
    return path


HANDOFF_PAGES = [
    ("Handoff Objective and Current Status", "This manual exists so another Codex instance can start product work without needing the whole chat. The current build is a Windows Node host plus phone PWA with local Wi-Fi control, host approval, real input, WebRTC high-quality stream, packaged zip, and a known v76 fix for touchpad/zoom behavior.", ["Treat `C:\\RemoteDesktopControllerPlan\\remote-control-mvp` as the source workspace.", "The project must be moved to the user's computer, then into a private GitHub repository before broad parallel work.", "Known unresolved issue: multi-monitor switching can still drift or return to the wrong screen; do not hide this from the next worker."]),
    ("Core Architecture", "The host serves the phone PWA, host console, capture page, REST APIs, WebSocket control channel, and RTC signaling. The phone sends validated commands; the host maps them to Windows input through the input adapter after safety approval.", ["`src/server.js` owns HTTP, WebSocket, capture launch, host state, safety endpoints, and RTC room integration.", "`src/input-adapter.js` owns real/dry-run mouse, wheel, keyboard, chord, and text input.", "`public/app.js` owns phone UI, pairing, touchpad, zoom, keyboard, monitor selection, RTC receiver, and stream rendering.", "`public/capture.js` owns the low-latency capture browser path.", "`packaging/local-wifi` builds the easy-send zip."]),
    ("Run, Test, and Package Commands", "The baseline commands are intentionally simple. A new worker should run them before editing and after any meaningful change.", ["Run host: `$env:HOST_KEY='dev-host-key'; npm start`.", "Open host: `http://127.0.0.1:4317/host?key=dev-host-key`.", "Run core tests: `node --test test\\protocol.test.js test\\rtc-room.test.js test\\settings-store.test.js test\\session-store.test.js test\\coordinate-mapper.test.js test\\input-adapter.test.js test\\host-console.test.js test\\phone-connection-help.test.js`.", "Package: `powershell -NoProfile -ExecutionPolicy Bypass -File packaging\\local-wifi\\package-local-wifi.ps1 -PackageName OpenHostLink`.", "Smoke package on a free port before sending it to family testers."]),
    ("Transfer to User Computer", "The safest transfer is a GitHub repository plus release zip. If GitHub is not ready yet, transfer the whole folder excluding generated junk and verify on the user's PC.", ["Copy source folders: `src`, `public`, `packaging`, `scripts`, `test`, `docs`, root `README.md`, `package.json`, `package-lock.json`, and current `dist\\OpenHostLink.zip`.", "Exclude `node_modules`, most `output` smoke folders, logs, PID files, local data/trust stores, and any personal screenshots.", "On the new PC install Node 20+, run `npm ci`, then run the tests and package smoke.", "If running from zip, use `Start Remote Controller.bat` or `start-local-wifi.ps1`; if running source, use `npm start`."]),
    ("GitHub Setup Plan", "Create a private repository first. Public/open-source release can come after security, licensing, and packaging are cleaner.", ["Initialize Git in the project root only after `.gitignore` is created.", "First commit: source, docs, tests, packaging scripts, package manifests; no `node_modules`, logs, host keys, or generated smoke output.", "Tag the imported baseline as `v0.1-local-wifi-baseline-v76`.", "Create Issues from the 25-page plan; each issue names owner, branch, files, tests, acceptance evidence.", "Use pull requests for every Codex instance branch."]),
    ("Parallel Codex Rules", "Every Codex instance needs a narrow workstream. Broad instructions like 'make it better' cause collisions.", ["Assign one chapter, one branch, and one file ownership boundary.", "Do not change shared protocol/auth/capture metadata without an interface note.", "Do not remove tests to pass CI.", "Do not rewrite the phone UI and host protocol in the same branch unless the issue explicitly owns both.", "Final report must include changed files, tests, screenshots/artifacts, risks, and next issue."]),
    ("Known Issues and Product Risks", "Known issues are part of the handoff, not embarrassment. They prevent repeated blind fixes.", ["Multi-monitor switching remains the biggest current functional risk.", "Browser capture cannot be the final unattended paid remote strategy because screen capture permission is user-mediated [S1].", "The project currently relies on Chrome/browser capture for high-quality stream; native capture/encode is the likely paid-product path.", "Relay bandwidth can destroy subscription margin; Cloudflare is promising for early TURN but must be measured [S4].", "App Store review needs careful positioning as a generic remote desktop client [S7]."]),
    ("Prompt Pack for New Instance", "Paste this prompt into a fresh Codex instance after the repo exists: 'You are working on the Remote Controller project. Read docs/product_planning/Remote_Controller_Handoff_and_Transfer_Manual.docx and docs/product_planning/Remote_Controller_Product_Execution_Plan_25_Pages.docx first. Work only on issue <ID> and branch <branch>. Preserve local Wi-Fi mode, host approval, real-input safety, v76 touchpad relative behavior, and package reproducibility. Do not edit files outside your ownership without asking. Before final response, run the issue's tests and update docs if behavior changed.'", ["Give the instance its chapter/page assignment.", "Give it the exact repo branch and file ownership.", "Give it one acceptance gate.", "Require a concise final report with evidence."]),
]


def create_handoff_doc():
    doc = Document()
    set_doc_styles(doc, "Remote Controller Handoff Manual")
    add_title(doc, "Remote Controller Handoff and Transfer Manual", "A practical manual for moving the current build to the user's computer, GitHub, and multiple Codex instances without losing project knowledge.")
    for idx, (title, lead, items) in enumerate(HANDOFF_PAGES, start=1):
        doc.add_heading(f"Page {idx}: {title}", level=1)
        p(doc, lead, bold_label="Page purpose:")
        for item in items:
            bullet(doc, item)
        for section in handoff_deepening_sections(title):
            doc.add_heading(section["heading"], level=2)
            p(doc, section["body"])
            for item in section["bullets"]:
                bullet(doc, item)
        add_goal_box(doc, [
            "A new worker can understand this page without chat history.",
            "The page gives concrete commands, files, or rules instead of vague summaries.",
            "Any risk or prerequisite named here has an owner in the 25-page execution plan.",
        ])
        if idx < len(HANDOFF_PAGES):
            doc.add_page_break()
    add_source_table(doc)
    add_source_appendix_detail(doc)
    path = OUT / "Remote_Controller_Handoff_and_Transfer_Manual.docx"
    doc.save(path)
    return path


def handoff_deepening_sections(title):
    return [
        {
            "heading": "Exact operator actions",
            "body": f"For '{title}', the operator should follow these steps rather than improvising from memory.",
            "bullets": [
                "Read the whole page once before touching files, Git, GitHub, package artifacts, or another computer.",
                "Capture current evidence first: folder path, branch or no-branch state, package version, latest zip path, and the commands that currently pass.",
                "Perform the named transfer or handoff action in one small batch, then verify immediately before moving to the next batch.",
                "Record any deviation in the handoff notes so the next Codex instance does not assume the ideal path happened.",
            ],
        },
        {
            "heading": "Step-by-step handoff script",
            "body": f"This is the minimum practical script for '{title}'. The point is to make the transfer repeatable for a new machine or a new Codex instance, not just understandable to the person who already lived through the prototype.",
            "bullets": [
                "Step 1: confirm the current workspace path and latest package path, then write both paths in the issue or handoff note before copying or editing anything.",
                "Step 2: run the listed smoke checks from the current machine so the sender knows whether the outgoing build is already healthy or already broken.",
                "Step 3: copy or publish only the approved source/package set. Exclude local runtime data, logs, PID files, extracted smoke folders, `node_modules`, and disabled scratch files.",
                "Step 4: on the receiving machine, launch from the documented command or the single packaged launcher and verify that the host page, phone URL, QR code, PIN, pairing, and real-input toggle appear as expected.",
                "Step 5: run the core automated tests from the receiving source checkout, or if using only the family package, run the package smoke and one phone pairing test.",
                "Step 6: record the result with machine name, Windows version, browser version, phone type, network type, and any missing permissions or firewall prompts.",
            ],
        },
        {
            "heading": "What must be true before the next person starts",
            "body": "A new worker should not begin creative implementation until the handoff basics are proven. This prevents debugging a broken transfer while pretending to build product features.",
            "bullets": [
                "The project opens from the documented folder and the expected files are present with normal names.",
                "The host can start or the package launcher can start without missing Node, missing modules, blocked scripts, or stale disabled files.",
                "The phone-facing page can be reached through the documented local IP or tunnel path and shows the current UI rather than an old cached version.",
                "The worker knows whether they are editing source, testing a package, creating a GitHub issue, or preparing a release artifact.",
                "The worker knows which known issues are accepted for now, especially monitor switching, native capture, remote relay economics, and app-store constraints.",
            ],
        },
        {
            "heading": "Files and state to protect",
            "body": "These are the project assets that should not be casually deleted, renamed, or overwritten during the transfer.",
            "bullets": [
                "`src/server.js`, `src/input-adapter.js`, `src/rtc-room.js`, `src/session-store.js`, `public/app.js`, `public/capture.js`, `public/host.js`, `public/sw.js`, and `packaging/local-wifi/*`.",
                "`docs/product_planning/*`, `README.md`, `MILESTONE_STATUS.md`, existing traceability docs, and checkpoint artifacts under `output/checkpoints`.",
                "`dist/OpenHostLink.zip` as the latest family-testable package, while treating it as a release artifact rather than source.",
                "Local secrets and runtime data under `data`, logs, PID files, and smoke-test extraction folders must be excluded from GitHub unless sanitized.",
            ],
        },
        {
            "heading": "Verification before handoff is accepted",
            "body": "A handoff is not complete because files were copied. It is complete when the recipient can run and verify the system.",
            "bullets": [
                "The recipient can install dependencies or use the packaged zip without missing-file errors.",
                "The host console opens, shows the correct version, and exposes a phone URL/QR code.",
                "The test command listed in this manual passes or has an explicit, documented reason why a physical-device gate is still pending.",
                "The next Codex instance has a branch, issue, file ownership boundary, and acceptance gate before it starts editing.",
            ],
        },
        {
            "heading": "Common mistakes to avoid",
            "body": "The transfer should avoid mistakes that already caused confusion during the prototype phase.",
            "bullets": [
                "Do not ship a package with multiple confusing launchers when one obvious launcher is expected.",
                "Do not let `.disabled-by-incident-response` file names become the canonical GitHub source state.",
                "Do not describe the monitor-switching system as fixed until it passes multi-monitor testing on another computer.",
                "Do not let a paid remote-access experiment break the free local Wi-Fi mode, because the free mode is the project's credibility base.",
            ],
        },
        {
            "heading": "Acceptance evidence to attach",
            "body": "Every handoff page should produce evidence. If the page does not create evidence, it is probably only a note and not a real manual step.",
            "bullets": [
                "Attach command output or a short transcript for every command that matters.",
                "Attach a screenshot when the result is visual, such as host console, QR code, package folder, phone UI, capture window, settings sheet, or display selector.",
                "Attach file hashes for release zips so the sender and receiver can prove they are testing the same artifact.",
                "Attach a short failed-test note when something is known broken. A truthful failed-test note is more useful than a vague 'needs polish' sentence.",
                "Attach the next action owner so the handoff does not end in a pile of unassigned concerns.",
            ],
        },
    ]


if __name__ == "__main__":
    handoff = create_handoff_doc()
    plan = create_plan_doc()
    print(handoff)
    print(plan)
